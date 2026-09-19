from __future__ import annotations

import html
import re
import urllib.error
import urllib.request
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.shared import Inches, Pt, RGBColor


BASE = Path(__file__).parent
SOURCE = BASE / "REPORT.md"
HTML_OUT = BASE / "REPORT.html"
DOCX_OUT = BASE / "REPORT.docx"
AUDIT_OUT = BASE / "source-link-audit.tsv"


def plain(text: str) -> str:
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    text = text.replace("**", "").replace("`", "")
    return text


def inline_html(text: str) -> str:
    escaped = html.escape(text)
    escaped = re.sub(
        r"\[([^\]]+)\]\(([^)]+)\)",
        r'<a href="\2">\1</a>',
        escaped,
    )
    escaped = re.sub(r"\*\*([^*]+)\*\*", r"<strong>\1</strong>", escaped)
    escaped = escaped.replace("`", "")
    return escaped


def is_table_separator(line: str) -> bool:
    return bool(re.fullmatch(r"\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)+\|?", line))


def cells(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def to_html(lines: list[str]) -> str:
    body: list[str] = []
    idx = 0
    list_open = False
    code_open = False
    while idx < len(lines):
        line = lines[idx].rstrip()
        if line.startswith("```"):
            if list_open:
                body.append("</ul>")
                list_open = False
            body.append("<pre>" if not code_open else "</pre>")
            code_open = not code_open
        elif code_open:
            body.append(html.escape(line))
        elif idx + 1 < len(lines) and line.startswith("|") and is_table_separator(lines[idx + 1]):
            if list_open:
                body.append("</ul>")
                list_open = False
            headers = cells(line)
            body.append("<table><thead><tr>" + "".join(f"<th>{inline_html(c)}</th>" for c in headers) + "</tr></thead><tbody>")
            idx += 2
            while idx < len(lines) and lines[idx].startswith("|"):
                row = cells(lines[idx])
                body.append("<tr>" + "".join(f"<td>{inline_html(c)}</td>" for c in row) + "</tr>")
                idx += 1
            body.append("</tbody></table>")
            continue
        elif line.startswith("### "):
            if list_open:
                body.append("</ul>")
                list_open = False
            body.append(f"<h3>{inline_html(line[4:])}</h3>")
        elif line.startswith("## "):
            if list_open:
                body.append("</ul>")
                list_open = False
            heading = line[3:]
            heading_class = " class='sources'" if heading == "핵심 출처" else ""
            body.append(f"<h2{heading_class}>{inline_html(heading)}</h2>")
        elif line.startswith("# "):
            if list_open:
                body.append("</ul>")
                list_open = False
            body.append(f"<h1>{inline_html(line[2:])}</h1>")
        elif re.match(r"^\d+\. ", line):
            if list_open:
                body.append("</ul>")
                list_open = False
            body.append(f"<p class='numbered'>{inline_html(line)}</p>")
        elif line.startswith("- "):
            if not list_open:
                body.append("<ul>")
                list_open = True
            body.append(f"<li>{inline_html(line[2:])}</li>")
        elif not line.strip():
            if list_open:
                body.append("</ul>")
                list_open = False
        else:
            if list_open:
                body.append("</ul>")
                list_open = False
            body.append(f"<p>{inline_html(line)}</p>")
        idx += 1
    if list_open:
        body.append("</ul>")
    return """<!doctype html><html lang='ko'><head><meta charset='utf-8'>
<style>
@page { size: A4; margin: 18mm 16mm 18mm 16mm; }
body { font-family: 'Noto Sans KR', 'Noto Sans', sans-serif; color:rgb(23,32,51); font-size:10.2pt; line-height:1.68; }
h1 { font-size:24pt; line-height:1.25; margin:0 0 16pt; color:rgb(20,37,61); }
h2 { font-size:16pt; line-height:1.35; margin:25pt 0 10pt; padding-top:5pt; border-top:1px solid rgb(185,198,212); color:rgb(18,78,120); page-break-after:avoid; }
h3 { font-size:12pt; margin:18pt 0 7pt; color:rgb(30,99,143); page-break-after:avoid; }
p { margin:0 0 8pt; } ul { margin:3pt 0 9pt; padding-left:20pt; } li { margin:2pt 0; }
table { border-collapse:collapse; width:100%; margin:10pt 0 13pt; font-size:8.2pt; page-break-inside:avoid; }
th { background:rgb(23,76,115); color:white; font-weight:700; } td, th { border:1px solid rgb(201,213,223); padding:5pt; vertical-align:top; word-break:keep-all; overflow-wrap:break-word; }
tr:nth-child(even) td { background:rgb(244,248,251); } a { color:rgb(7,93,153); text-decoration:none; } pre { white-space:pre-wrap; background:rgb(244,247,250); padding:10pt; font-size:8.6pt; }
.numbered { margin:10pt 0 8pt 10pt; } li { break-inside:avoid; }
h2.sources + ul { font-size:8.1pt; line-height:1.3; }
h2.sources ~ p.numbered { font-size:7.7pt; line-height:1.2; margin:2pt 0 2pt 10pt; }
</style></head><body>""" + "\n".join(body) + "</body></html>"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_row_pagination(row, header: bool = False) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)
    if header:
        table_header = OxmlElement("w:tblHeader")
        table_header.set(qn("w:val"), "true")
        tr_pr.append(table_header)


def style_run(run, bold: bool = False) -> None:
    run.bold = bold
    run.font.name = "Noto Sans KR"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans KR")
    run.font.size = Pt(9.5)


def add_hyperlink(paragraph, label: str, url: str) -> None:
    relation_id = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relation_id)
    run = OxmlElement("w:r")
    run_props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "075D99")
    run_props.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_props.append(underline)
    east_asia = OxmlElement("w:rFonts")
    east_asia.set(qn("w:eastAsia"), "Noto Sans KR")
    run_props.append(east_asia)
    run.append(run_props)
    text_node = OxmlElement("w:t")
    text_node.text = label
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_text(paragraph, text: str, bold: bool = False) -> None:
    cursor = 0
    for match in re.finditer(r"\[([^\]]+)\]\((https?://[^)]+)\)", text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor:match.start()].replace("**", "").replace("`", ""))
            style_run(run, bold)
        add_hyperlink(paragraph, match.group(1), match.group(2))
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:].replace("**", "").replace("`", ""))
        style_run(run, bold)


def to_docx(lines: list[str]) -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)
    styles = doc.styles
    styles["Normal"].font.name = "Noto Sans KR"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans KR")
    styles["Normal"].font.size = Pt(9.5)
    for level, size, color in ((1, 20, "14253D"), (2, 14, "124E78"), (3, 11, "1E638F")):
        style = styles[f"Heading {level}"]
        style.font.name = "Noto Sans KR"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans KR")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_text(header, "폰티스형 선교·영리 생태계 사례 조사")
    idx = 0
    while idx < len(lines):
        line = lines[idx].rstrip()
        if idx + 1 < len(lines) and line.startswith("|") and is_table_separator(lines[idx + 1]):
            header_cells = cells(line)
            idx += 2
            rows: list[list[str]] = []
            while idx < len(lines) and lines[idx].startswith("|"):
                rows.append(cells(lines[idx]))
                idx += 1
            table = doc.add_table(rows=1, cols=len(header_cells))
            table.style = "Table Grid"
            set_row_pagination(table.rows[0], header=True)
            for col, value in enumerate(header_cells):
                cell = table.rows[0].cells[col]
                set_cell_shading(cell, "174C73")
                add_text(cell.paragraphs[0], value, True)
            for row_values in rows:
                row = table.add_row().cells
                set_row_pagination(table.rows[-1])
                for col, value in enumerate(row_values):
                    add_text(row[col].paragraphs[0], value)
            continue
        if line.startswith("# "):
            doc.add_heading(plain(line[2:]), level=1)
        elif line.startswith("## "):
            doc.add_heading(plain(line[3:]), level=2)
        elif line.startswith("### "):
            doc.add_heading(plain(line[4:]), level=3)
        elif line.startswith("- "):
            paragraph = doc.add_paragraph(style="List Bullet")
            add_text(paragraph, line[2:])
        elif re.match(r"^\d+\. ", line):
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.12)
            paragraph.paragraph_format.first_line_indent = Inches(-0.12)
            add_text(paragraph, line)
        elif line.startswith("```"):
            pass
        elif line.strip():
            paragraph = doc.add_paragraph()
            add_text(paragraph, line)
        idx += 1
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(footer, "공개 자료 조사 · 법률 자문 아님")
    doc.save(DOCX_OUT)


def check_url(url: str) -> tuple[str, str]:
    request = urllib.request.Request(url, method="HEAD", headers={"User-Agent": "Mozilla/5.0 research-audit"})
    try:
        with urllib.request.urlopen(request, timeout=15) as response:
            return url, f"open:{response.status}"
    except urllib.error.HTTPError as exc:
        if exc.code in (403, 405):
            try:
                request = urllib.request.Request(url, method="GET", headers={"User-Agent": "Mozilla/5.0 research-audit", "Range": "bytes=0-1024"})
                with urllib.request.urlopen(request, timeout=15) as response:
                    return url, f"open-get:{response.status}"
            except Exception as fallback_exc:
                return url, f"inaccessible:{type(fallback_exc).__name__}"
        return url, f"http:{exc.code}"
    except Exception as exc:
        return url, f"inaccessible:{type(exc).__name__}"


def audit_sources(markdown: str) -> None:
    urls = sorted(set(re.findall(r"\]\((https?://[^)]+)\)", markdown)))
    with ThreadPoolExecutor(max_workers=10) as executor:
        results = list(executor.map(check_url, urls))
    AUDIT_OUT.write_text("url\tstatus\n" + "\n".join(f"{url}\t{status}" for url, status in results) + "\n", encoding="utf-8")


def main() -> None:
    markdown = SOURCE.read_text(encoding="utf-8")
    lines = markdown.splitlines()
    HTML_OUT.write_text(to_html(lines), encoding="utf-8")
    to_docx(lines)
    audit_sources(markdown)
    print(HTML_OUT)
    print(DOCX_OUT)
    print(AUDIT_OUT)


if __name__ == "__main__":
    main()
